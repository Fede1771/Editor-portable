"""Conversión de documentos PDF a otros formatos (fase 6): Word, Excel, Imágenes, Texto, HTML."""
from __future__ import annotations

import os
from typing import Sequence

import pymupdf as fitz

from .attachments import unique_path


def _safe_excel_cell(val: object | None) -> str:
    """Normaliza un valor extraído a cadena literal para Excel sin insertar caracteres espurios (H20, N05)."""
    if val is None:
        return ""
    return str(val)


def _append_literal_row(ws, values: Sequence[object | None], next_row_idx: int) -> int:
    """Escribe una fila en `ws` forzando tipo de celda texto (`data_type='s'`) sin evaluar fórmulas
    ni alterar el valor literal con apóstrofos iniciales (H20, N05)."""
    from openpyxl.cell.cell import TYPE_STRING

    for col_idx, val in enumerate(values, start=1):
        s = _safe_excel_cell(val)
        cell = ws.cell(row=next_row_idx, column=col_idx)
        cell._value = s
        cell.data_type = TYPE_STRING
    return next_row_idx + 1


def convert_to_word(doc: fitz.Document, out_docx_path: str, src_path: str | None = None) -> None:
    """Convierte el documento PDF en vivo a Microsoft Word (.docx) (H12, H26)."""
    from pdf2docx import Converter
    from .privacy import TempManager

    fd, temp_pdf = TempManager.create_temp_file(prefix="pdf_word_", suffix=".pdf")
    os.close(fd)
    doc.save(temp_pdf, garbage=3, deflate=True)
    try:
        cv = Converter(temp_pdf)
        cv.convert(out_docx_path)
        cv.close()
    finally:
        TempManager.cleanup_path(temp_pdf)


def convert_to_excel(doc: fitz.Document, out_xlsx_path: str) -> int:
    """Extrae tablas del PDF y las guarda en un libro de Microsoft Excel (.xlsx) como texto literal (H20, N05).

    Retorna la cantidad total de tablas extraídas.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    default_sheet = wb.active

    total_tables = 0
    for pno in range(doc.page_count):
        page = doc[pno]
        tabs = page.find_tables()
        for t_idx, table in enumerate(tabs):
            total_tables += 1
            sheet_title = f"Pág {pno + 1} Tab {t_idx + 1}"[:31]
            ws = wb.create_sheet(title=sheet_title)
            row_idx = 1
            for row in table.extract():
                row_idx = _append_literal_row(ws, row, row_idx)

    # Si no se detectaron tablas tabuladas estructuradas, exportar líneas de texto por renglón
    if total_tables == 0:
        ws = wb.create_sheet(title="Contenido Texto")
        row_idx = 1
        for pno in range(doc.page_count):
            row_idx = _append_literal_row(ws, [f"--- PÁGINA {pno + 1} ---"], row_idx)
            for line in doc[pno].get_text("text").splitlines():
                if line.strip():
                    row_idx = _append_literal_row(ws, [line.strip()], row_idx)

    if len(wb.sheetnames) > 1 and default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    os.makedirs(os.path.dirname(os.path.abspath(out_xlsx_path)), exist_ok=True)
    wb.save(out_xlsx_path)
    return total_tables


def convert_to_images(
    doc: fitz.Document,
    out_dir: str,
    fmt: str = "png",
    dpi: int = 200,
    pages: Sequence[int] | None = None,
) -> list[str]:
    """Exporta las páginas del PDF como imágenes de alta resolución sin sobrescribir archivos existentes (H28)."""
    os.makedirs(out_dir, exist_ok=True)
    target_pages = list(pages) if pages is not None else list(range(doc.page_count))
    exported: list[str] = []

    ext = fmt.lower().replace(".", "")
    for pno in target_pages:
        page = doc[pno]
        pix = page.get_pixmap(dpi=dpi)
        filename = f"pagina_{pno + 1:03d}.{ext}"
        out_path = unique_path(os.path.join(out_dir, filename))
        pix.save(out_path)
        exported.append(out_path)

    return exported


def convert_to_text(doc: fitz.Document, out_txt_path: str) -> None:
    """Exporta el texto plano de todo el documento."""
    os.makedirs(os.path.dirname(os.path.abspath(out_txt_path)), exist_ok=True)
    with open(out_txt_path, "w", encoding="utf-8") as f:
        for pno in range(doc.page_count):
            f.write(f"--- PÁGINA {pno + 1} ---\n\n")
            f.write(doc[pno].get_text("text"))
            f.write("\n\n")


def convert_to_html(doc: fitz.Document, out_html_path: str) -> None:
    """Exporta el contenido del documento a formato HTML con estilos."""
    os.makedirs(os.path.dirname(os.path.abspath(out_html_path)), exist_ok=True)
    with open(out_html_path, "w", encoding="utf-8") as f:
        f.write("<!DOCTYPE html><html><head><meta charset='utf-8'><title>Documento Exportado</title>")
        f.write("<style>body{font-family:sans-serif;margin:40px;background:#f0f2f5;}.page{background:#fff;padding:30px;margin-bottom:30px;box-shadow:0 2px 6px rgba(0,0,0,0.15);}</style>")
        f.write("</head><body>\n")
        for pno in range(doc.page_count):
            f.write(f"<div class='page'><h3>Página {pno + 1}</h3>\n")
            f.write(doc[pno].get_text("html"))
            f.write("</div>\n")
        f.write("</body></html>")
