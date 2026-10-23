"""Integración Office: rutas, atomicidad, datos en vivo y límites de proceso."""
from pathlib import Path
import time

import openpyxl
import pymupdf as fitz
import pytest

from app import conversion, creation, office365


def _stuck_worker(conn, *args):
    time.sleep(20)


@pytest.mark.parametrize("extension,operation", [("docx", "word_pdf"), ("xlsx", "excel_pdf"), ("pptx", "powerpoint_pdf")])
def test_m365_office_to_pdf_routes_native_application(tmp_path, monkeypatch, extension, operation):
    source = tmp_path / ("original." + extension)
    source.write_bytes(b"UNCHANGED")
    def export(kind, input_path, output_path, **kwargs):
        assert kind == operation and input_path == str(source)
        with fitz.open() as doc:
            doc.new_page().insert_text((30, 50), "NATIVE_OFFICE")
            doc.save(output_path)
    monkeypatch.setattr(office365, "run_office", export)
    out = creation.office_to_pdf_file(str(source), str(tmp_path / "out"))
    assert source.read_bytes() == b"UNCHANGED"
    with fitz.open(out) as doc:
        assert "NATIVE_OFFICE" in doc[0].get_text()


def test_m365_failed_office_export_keeps_previous_pdf(tmp_path, monkeypatch):
    source = tmp_path / "original.docx"
    source.write_bytes(b"UNCHANGED")
    output = tmp_path / "original.pdf"
    with fitz.open() as doc:
        doc.new_page().insert_text((30, 50), "OLD_VALID_PDF")
        doc.save(output)
    before = output.read_bytes()
    def invalid(*args, **kwargs):
        Path(args[2]).write_bytes(b"CORRUPT")
    monkeypatch.setattr(office365, "run_office", invalid)
    with pytest.raises(Exception):
        creation.office_to_pdf_file(str(source), str(tmp_path))
    assert output.read_bytes() == before
    assert not list(tmp_path.glob("*.tmp.pdf"))


def test_m365_excel_uses_current_unsaved_pdf_and_literal_values(tmp_path, monkeypatch):
    with fitz.open() as doc:
        doc.new_page().insert_text((30, 50), "UNSAVED_EDIT")
        captured = []
        def native(kind, source, target):
            assert kind == "pdf_excel"
            captured.append(source)
            with fitz.open(source) as snapshot:
                assert "UNSAVED_EDIT" in snapshot[0].get_text()
            book = openpyxl.Workbook()
            book.active["A1"] = "=1+1"
            book.active["A1"].data_type = "s"
            book.save(target)
            book.close()
            return 1
        monkeypatch.setattr(office365, "run_office", native)
        output = tmp_path / "literal.xlsx"
        assert conversion.convert_to_excel(doc, str(output)) == 1
    assert not Path(captured[0]).exists()
    book = openpyxl.load_workbook(output)
    assert book.active["A1"].value == "=1+1" and book.active["A1"].data_type == "s"
    book.close()


def test_m365_invalid_formula_output_does_not_replace_previous_file(tmp_path, monkeypatch):
    output = tmp_path / "report.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = "OLD"
    book.save(output)
    book.close()
    before = output.read_bytes()
    def unsafe(kind, source, target):
        book = openpyxl.Workbook()
        book.active["A1"] = "=1+1"
        book.save(target)
        book.close()
    monkeypatch.setattr(office365, "run_office", unsafe)
    with fitz.open() as doc:
        doc.new_page()
        with pytest.raises(ValueError, match="fórmulas"):
            conversion.convert_to_excel(doc, str(output))
    assert output.read_bytes() == before


def test_m365_timeout_terminates_only_worker(tmp_path, monkeypatch):
    source = tmp_path / "input.docx"
    source.write_bytes(b"DUMMY")
    monkeypatch.setattr(office365, "_office_process", _stuck_worker)
    before = time.monotonic()
    with pytest.raises(office365.OfficeConversionError, match="no completó"):
        office365.run_office("word_pdf", str(source), str(tmp_path / "out.pdf"), timeout=.2)
    assert time.monotonic() - before < 8


def test_m365_cleanup_rejects_reused_pid(monkeypatch):
    class Unrelated:
        def create_time(self):
            return 99
        def name(self):
            return "WINWORD.EXE"
        def kill(self):
            pytest.fail("No terminar un PID reutilizado")
    monkeypatch.setattr(office365.psutil, "Process", lambda pid: Unrelated())
    office365._stop_owned({"pid": 123, "created": 1, "name": "winword.exe"})


def test_m365_m_literal_escapes_path_without_code_injection():
    assert office365._m_string('a"#(lf)b') == '"a""#(#)(lf)b"'


def test_m365_existing_powerpoint_is_not_opened_or_closed(monkeypatch):
    import win32com.client
    class Existing:
        pid = 123
        info = {"name": "POWERPNT.EXE"}
    monkeypatch.setattr(office365.psutil, "process_iter", lambda *args: [Existing()])
    def unexpected(*args):
        pytest.fail("No conectar a PowerPoint del usuario")
    monkeypatch.setattr(win32com.client, "DispatchEx", unexpected)
    with pytest.raises(office365.OfficeConversionError, match="PowerPoint está abierto"):
        with office365._application("PowerPoint", None):
            pytest.fail("No abrir ni modificar una presentación del usuario")


def test_m365_parallel_jobs_cannot_share_application_lock():
    from concurrent.futures import ThreadPoolExecutor
    import uuid
    name = "Test_" + uuid.uuid4().hex
    def second():
        with pytest.raises(office365.OfficeConversionError, match="Otra conversión"):
            with office365._office_lock(name, .05):
                pytest.fail("Dos trabajos no deben adquirir el mismo bloqueo")
    with office365._office_lock(name, 1):
        with ThreadPoolExecutor(max_workers=1) as executor:
            executor.submit(second).result(timeout=3)
