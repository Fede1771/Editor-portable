"""Conversión de documentos PDF a otros formatos (fase 6): Word, Excel, Imágenes, Texto, HTML."""
from __future__ import annotations

import os
import tempfile
import zipfile
from typing import Callable, Sequence

import pymupdf as fitz

from .attachments import unique_path


def _atomic_write_target(
    out_path: str,
    suffix: str,
    writer_fn: Callable[[str], None],
    verifier_fn: Callable[[str], None] | None = None,
) -> None:
    """Escribe en un archivo temporal en el mismo directorio, verifica y reemplaza atómicamente (V07)."""
    target = os.path.abspath(out_path)
    folder = os.path.dirname(target) or "."
    os.makedirs(folder, exist_ok=True)
    fd, temp_path = tempfile.mkstemp(
        prefix=f".{os.path.basename(target)}.",
        suffix=suffix,
        dir=folder,
    )
    os.close(fd)
    try:
        writer_fn(temp_path)
        if not os.path.isfile(temp_path) or os.path.getsize(temp_path) == 0:
            raise OSError("El archivo de exportación generado está vacío o incompleto.")
        if verifier_fn is not None:
            verifier_fn(temp_path)
        os.replace(temp_path, target)
    finally:
        if os.path.exists(temp_path):
            from .privacy import secure_wipe_file
            secure_wipe_file(temp_path)


def _safe_excel_cell(val: object | None) -> str:
    """Normaliza un valor extraído a cadena literal para Excel sin insertar caracteres espurios (H20, N05)."""
    if val is None:
        return ""
    return str(val)


def _append_literal_row(ws, values: Sequence[object | None], next_row_idx: int) -> int:
    """Escribe una fila en `ws` forzando tipo de celda texto (`data_type='s'`) sin evaluar fórmulas
    ni alterar el valor literal con apóstrofos iniciales (H20, N05)."""
    from openpyxl.cell.cell import TYPE_STRING

    for col_idx, val in enumerate(values, start=1):
        s = _safe_excel_cell(val)
        cell = ws.cell(row=next_row_idx, column=col_idx)
        cell._value = s
        cell.data_type = TYPE_STRING
    return next_row_idx + 1


def _verify_docx(path: str) -> None:
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None or not {"[Content_Types].xml", "word/document.xml"} <= set(archive.namelist()):
            raise ValueError("Microsoft Word no generó un documento OOXML válido.")


def _verify_xlsx(path: str) -> None:
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if not book.sheetnames:
            raise ValueError("Microsoft Excel no generó hojas válidas.")
        if any(cell.data_type == "f" for sheet in book for row in sheet for cell in row):
            raise ValueError("La exportación del PDF debe contener texto literal, no fórmulas.")
    finally:
        book.close()


def _office_export(doc, target, operation, suffix, verifier):
    from .office365 import run_office
    from .privacy import TempManager
    fd, temp_pdf = TempManager.create_temp_file(prefix="pdf_office_", suffix=".pdf")
    os.close(fd)
    result = None
    intermediate = None
    try:
        # Copia única del documento en vivo; no se abre ni modifica el original.
        doc.save(temp_pdf, garbage=3, deflate=True)
        source = temp_pdf
        if operation == "word_docx":
            from pdf2docx import Converter
            intermediate = TempManager.create_temp_dir(prefix="word_structure_")
            source = os.path.join(intermediate, "structure.docx")
            converter = Converter(temp_pdf)
            try:
                converter.convert(source)
            finally:
                converter.close()
            _verify_docx(source)
        def write(output):
            nonlocal result
            result = run_office(operation, source, output)
        _atomic_write_target(target, suffix, write, verifier)
        return result
    finally:
        if intermediate is not None:
            TempManager.cleanup_path(intermediate)
        TempManager.cleanup_path(temp_pdf)


def convert_to_word(doc: fitz.Document, out_docx_path: str, src_path: str | None = None) -> None:
    """Reconstruye el PDF y Microsoft Word de escritorio guarda el DOCX final verificado."""
    _office_export(doc, out_docx_path, "word_docx", ".tmp.docx", _verify_docx)


def convert_to_excel(doc: fitz.Document, out_xlsx_path: str) -> int:
    """Excel/Power Query importa el PDF; el XLSX final contiene valores de texto."""
    return _office_export(doc, out_xlsx_path, "pdf_excel", ".tmp.xlsx", _verify_xlsx)


def convert_to_images(
    doc: fitz.Document,
    out_dir: str,
    fmt: str = "png",
    dpi: int = 200,
    pages: Sequence[int] | None = None,
) -> list[str]:
    """Exporta las páginas del PDF como imágenes de alta resolución sin sobrescribir archivos existentes (H28, V08)."""
    from .viewer import render_bounded_pixmap

    os.makedirs(out_dir, exist_ok=True)
    target_pages = list(pages) if pages is not None else list(range(doc.page_count))
    exported: list[str] = []

    ext = fmt.lower().replace(".", "")
    for pno in target_pages:
        page = doc[pno]
        pix = render_bounded_pixmap(page, dpi=dpi, alpha=False)
        filename = f"pagina_{pno + 1:03d}.{ext}"
        out_path = unique_path(os.path.join(out_dir, filename))
        pix.save(out_path)
        exported.append(out_path)

    return exported


def convert_to_text(doc: fitz.Document, out_txt_path: str) -> None:
    """Exporta el texto plano de todo el documento de forma atómica (V07)."""
    def _write_txt(temp_txt: str) -> None:
        with open(temp_txt, "w", encoding="utf-8") as f:
            for pno in range(doc.page_count):
                f.write(f"--- PÁGINA {pno + 1} ---\n\n")
                f.write(doc[pno].get_text("text"))
                f.write("\n\n")

    _atomic_write_target(out_txt_path, ".tmp.txt", _write_txt)


def convert_to_html(doc: fitz.Document, out_html_path: str) -> None:
    """Exporta el contenido del documento a formato HTML con estilos de forma atómica (V07)."""
    def _write_html(temp_html: str) -> None:
        with open(temp_html, "w", encoding="utf-8") as f:
            f.write("<!DOCTYPE html><html><head><meta charset='utf-8'><title>Documento Exportado</title>")
            f.write("<style>body{font-family:sans-serif;margin:40px;background:#f0f2f5;}.page{background:#fff;padding:30px;margin-bottom:30px;box-shadow:0 2px 6px rgba(0,0,0,0.15);}</style>")
            f.write("</head><body>\n")
            for pno in range(doc.page_count):
                f.write(f"<div class='page'><h3>Página {pno + 1}</h3>\n")
                f.write(doc[pno].get_text("html"))
                f.write("</div>\n")
            f.write("</body></html>")

    _atomic_write_target(out_html_path, ".tmp.html", _write_html)
